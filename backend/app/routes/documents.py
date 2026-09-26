from ..workflows.queue import submit_job
import io
from fastapi import Depends, HTTPException, Header
from fastapi.responses import StreamingResponse
from docx import Document
from ..security import current_user, db_session
from ..schemas import DocumentEdit, DocumentRequest
from ..store import owned
from ..documents import approve_edit

from fastapi import APIRouter
from ..http_common import PREFIX, request_key, public_record

router = APIRouter()

@router.post(PREFIX + '/documents')
def create_document(body: DocumentRequest, user=Depends(current_user), db=Depends(db_session), idempotency_key: str | None = Header(None)):
    owned(db, body.cv_id, user.id, 'cv')
    owned(db, body.vacancy_id, user.id, 'vacancy')
    return submit_job(db, user.id, 'document', body.model_dump(), request_key(idempotency_key))


@router.put(PREFIX + '/documents/{document_id}')
def edit_document(document_id: str, body: DocumentEdit, user=Depends(current_user), db=Depends(db_session), if_match: str | None = Header(None)):
    row = owned(db, document_id, user.id, 'document', lock=True)
    if if_match is not None and if_match != row.updated_at.isoformat():
        raise HTTPException(409, 'Объект изменён. Обновите данные перед сохранением.')
    approve_edit(row, body)
    db.commit()
    return public_record(row)


@router.get(PREFIX + '/documents/{document_id}/export')
def export_document(document_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, document_id, user.id, 'document')
    if row.data.get('requires_confirmation') and not row.data.get('approved'):
        raise HTTPException(409, 'Сначала проверьте и сохраните документ.')
    document = Document()
    document.add_heading(row.data['title'], 0)
    for paragraph in row.data['text'].split('\n'):
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                             headers={'Content-Disposition': 'attachment; filename="jobfinder-document.docx"'})


