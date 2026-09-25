import hashlib
import io
import secrets
import shutil
import threading
import time
from collections import defaultdict, deque
from pathlib import Path
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File, Header
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import select, text, delete
from sqlalchemy.exc import IntegrityError
from argon2.exceptions import VerificationError
from docx import Document
from pydantic import BaseModel, Field
from .config import settings
from .db import Session, User, Login, Record, Job, Knowledge, Budget, Usage, uid, now
from .security import current_user, admin, db_session, passwords, digest, new_session, user_dict
from .schemas import *
from .store import owned, shared, record_dict, save_data, enqueue
from .ingest import youtube_id, extract_cv
from .retrieval import retrieve, evidence
from .cv_sections import section_draft
from .ai import text_available

app = FastAPI(title='JobFinderKZ', version='0.1.0', docs_url='/api/docs')
app.add_middleware(CORSMiddleware, allow_origins=[settings.app_origin], allow_credentials=True,
                   allow_methods=['GET', 'POST', 'PUT', 'PATCH', 'DELETE'], allow_headers=['Content-Type', 'X-CSRF-Token', 'Idempotency-Key'])
PREFIX = '/api/v1'
attempts = defaultdict(deque)
attempts_lock = threading.Lock()


@app.middleware('http')
async def security_headers(request, call_next):
    origin = request.headers.get('origin')
    if request.method not in ('GET', 'HEAD', 'OPTIONS') and origin and origin != settings.app_origin:
        return Response('Недопустимый Origin', status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Cache-Control'] = 'no-store'
    return response


@app.exception_handler(ValueError)
async def invalid_value(request, exc):
    from fastapi.responses import JSONResponse
    return JSONResponse(status_code=422, content={'detail': str(exc)})


def throttle(request):
    host = request.client.host
    with attempts_lock:
        ticks = attempts[host]
        current = time.monotonic()
        while ticks and ticks[0] < current - 60:
            ticks.popleft()
        if len(ticks) >= 15:
            raise HTTPException(429, 'Слишком много попыток. Подождите минуту.')
        ticks.append(current)


def request_key(value):
    if value and (len(value) > 100 or not value.isascii()):
        raise HTTPException(422, 'Неверный Idempotency-Key')
    return value or uid()


def public_record(row):
    value = record_dict(row)
    value['data'] = {k: v for k, v in value['data'].items() if k != 'file_path'}
    if row.kind == 'cv' and not value['data'].get('facts'):
        value['data'] = {**value['data'], **section_draft(value['data'].get('text', ''))}
    return value


@app.get(PREFIX + '/health')
def health(db=Depends(db_session)):
    db.execute(text('SELECT 1'))
    return {'status': 'ok'}


@app.post(PREFIX + '/auth/register')
def register(body: Credentials, request: Request, response: Response, db=Depends(db_session)):
    throttle(request)
    user = User(email=body.email, password_hash=passwords.hash(body.password), profile=Profile().model_dump(),
                role='admin' if body.email == settings.admin_email.lower() else 'user')
    db.add(user)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, 'Этот email уже зарегистрирован')
    return new_session(db, user, response)


@app.post(PREFIX + '/auth/login')
def login(body: Credentials, request: Request, response: Response, db=Depends(db_session)):
    throttle(request)
    user = db.scalar(select(User).where(User.email == body.email))
    try:
        if not user or not passwords.verify(user.password_hash, body.password):
            raise HTTPException(401, 'Неверный email или пароль')
    except VerificationError:
        raise HTTPException(401, 'Неверный email или пароль')
    return new_session(db, user, response)


@app.get(PREFIX + '/auth/me')
def me(request: Request, user=Depends(current_user)):
    return {'user': user_dict(user), 'csrf': request.state.csrf}


@app.post(PREFIX + '/auth/logout')
def logout(request: Request, response: Response, user=Depends(current_user), db=Depends(db_session)):
    db.execute(delete(Login).where(Login.token == digest(request.cookies.get('jf_session', ''))))
    db.commit()
    response.delete_cookie('jf_session', path='/')
    return {'ok': True}


@app.put(PREFIX + '/profile')
def profile(body: Profile, user=Depends(current_user), db=Depends(db_session)):
    user.profile = body.model_dump()
    db.commit()
    return user_dict(user)


@app.delete(PREFIX + '/account')
def delete_account(response: Response, user=Depends(current_user), db=Depends(db_session)):
    # Same key as worker: deletion cannot race a job that would recreate files.
    if not db.scalar(text('SELECT pg_try_advisory_xact_lock(hashtext(:key))'), {'key': 'user:' + user.id}):
        raise HTTPException(409, 'Дождитесь завершения текущего задания и повторите удаление')
    folder = (settings.storage_path / 'users' / user.id).resolve()
    root = (settings.storage_path / 'users').resolve()
    if folder.parent != root:
        raise HTTPException(500, 'Некорректный путь хранилища')
    shutil.rmtree(folder, ignore_errors=True)
    db.delete(user)
    db.commit()
    response.delete_cookie('jf_session', path='/')
    return {'ok': True}


@app.get(PREFIX + '/records/{kind}')
def records(kind: str, user=Depends(current_user), db=Depends(db_session)):
    if kind not in ('cv', 'vacancy', 'document', 'plan', 'interview'):
        raise HTTPException(404)
    rows = db.scalars(select(Record).where(Record.kind == kind, Record.owner_id == user.id).order_by(Record.created_at.desc())).all()
    return [public_record(row) for row in rows]


@app.get(PREFIX + '/record/{record_id}')
def record(record_id: str, user=Depends(current_user), db=Depends(db_session)):
    return public_record(owned(db, record_id, user.id))


async def upload(file, user_id, extensions, limit):
    suffix = Path(file.filename or '').suffix.lower()
    if suffix not in extensions:
        raise HTTPException(422, 'Неподдерживаемый формат файла')
    directory = settings.storage_path / 'users' / user_id
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / (uid() + suffix)
    size = 0
    try:
        with path.open('xb') as output:
            while block := await file.read(1024 * 1024):
                size += len(block)
                if size > limit:
                    raise HTTPException(413, 'Файл слишком большой')
                output.write(block)
        if not size:
            raise HTTPException(422, 'Файл пуст')
    except Exception:
        path.unlink(missing_ok=True)
        raise
    finally:
        await file.close()
    return path


@app.post(PREFIX + '/cv/text')
def cv_text(body: CVText, user=Depends(current_user), db=Depends(db_session)):
    row = Record(kind='cv', owner_id=user.id, status='review', data={'text': body.text, 'filename': 'Вставленный текст', **section_draft(body.text)})
    db.add(row)
    db.commit()
    return public_record(row)


@app.post(PREFIX + '/cv/upload')
async def cv_upload(file: UploadFile = File(...), user=Depends(current_user), db=Depends(db_session)):
    filename = Path(file.filename or 'CV').name
    path = await upload(file, user.id, {'.pdf', '.docx'}, 10_000_000)
    try:
        value = extract_cv(path)
    except Exception as exc:
        path.unlink(missing_ok=True)
        if isinstance(exc, ValueError):
            raise
        raise HTTPException(422, 'Не удалось прочитать документ. Проверьте PDF/DOCX или вставьте текст вручную.') from exc
    row = Record(kind='cv', owner_id=user.id, status='review', data={'text': value, 'file_path': str(path), 'filename': filename, **section_draft(value)})
    db.add(row)
    db.commit()
    return public_record(row)


@app.get(PREFIX + '/cv/{cv_id}/original')
def cv_original(cv_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, cv_id, user.id, 'cv')
    if not row.data.get('file_path'):
        raise HTTPException(404, 'Резюме добавлено текстом')
    return FileResponse(row.data['file_path'], filename=row.data['filename'])


@app.post(PREFIX + '/cv/{cv_id}/parse')
def cv_parse(cv_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, cv_id, user.id, 'cv')
    return enqueue(db, user.id, 'parse_cv', {'cv_id': row.id}, 'parse:' + cv_id)


@app.put(PREFIX + '/cv/{cv_id}/confirm')
def cv_confirm(cv_id: str, body: CVFacts, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, cv_id, user.id, 'cv')
    previous = row.data.get('facts')
    versions = row.data.get('versions', [])
    if previous:
        versions = versions + [{'facts': previous, 'saved_at': now().isoformat()}]
    save_data(row, facts=body.model_dump(), versions=versions)
    row.status = 'confirmed'
    db.commit()
    return public_record(row)


@app.post(PREFIX + '/cv/{cv_id}/sections')
def cv_sections(cv_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, cv_id, user.id, 'cv', lock=True)
    versions = row.data.get('versions', [])
    if row.data.get('facts'):
        versions = versions + [{'facts': row.data['facts'], 'saved_at': now().isoformat()}]
    save_data(row, **section_draft(row.data['text']), versions=versions)
    row.status = 'review'
    db.commit()
    return public_record(row)


@app.post(PREFIX + '/vacancies')
def vacancy_create(body: VacancyInput, user=Depends(admin), db=Depends(db_session)):
    data = body.model_dump()
    key = hashlib.sha256((body.url or body.title + body.description).strip().encode()).hexdigest()
    existing = db.scalar(select(Record).where(Record.owner_id == user.id, Record.kind == 'vacancy', Record.dedup_key == key))
    if existing:
        return public_record(existing)
    row = Record(kind='vacancy', owner_id=user.id, dedup_key=key, status='saved', data={**data, 'source': 'manual', 'favorite': False})
    db.add(row)
    db.commit()
    return public_record(row)


@app.post(PREFIX + '/vacancies/{vacancy_id}/favorite')
def favorite(vacancy_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, vacancy_id, user.id, 'vacancy', lock=True)
    save_data(row, favorite=not row.data.get('favorite', False))
    db.commit()
    return public_record(row)


class RankRequest(Strict):
    cv_id: str
    vacancy_ids: list[str] = Field(min_length=1, max_length=20)


@app.post(PREFIX + '/vacancies/rank')
def rank_create(body: RankRequest, user=Depends(current_user), db=Depends(db_session), idempotency_key: str | None = Header(None)):
    owned(db, body.cv_id, user.id, 'cv')
    for rid in body.vacancy_ids:
        owned(db, rid, user.id, 'vacancy')
    return enqueue(db, user.id, 'rank', body.model_dump(), request_key(idempotency_key))


@app.post(PREFIX + '/vacancies/hh/sync')
def hh_sync(body: HHQuery, user=Depends(current_user), db=Depends(db_session), idempotency_key: str | None = Header(None)):
    payload = body.model_dump()
    if body.area is None and body.regions is None:
        payload['regions'] = user.profile.get('regions', ['Казахстан'])
    return enqueue(db, user.id, 'hh_sync', payload, request_key(idempotency_key))


@app.get(PREFIX + '/connections')
def connections(user=Depends(current_user), db=Depends(db_session)):
    last = db.scalar(select(Job).where(Job.owner_id == user.id, Job.kind == 'hh_sync').order_by(Job.created_at.desc()))
    return {'openai': bool(settings.openai_api_key), 'text_ai': text_available(),
        'text_provider': settings.text_provider, 'gemini_free_tier': settings.text_provider == 'gemini' and settings.gemini_free_tier,
        'audio': bool(settings.openai_api_key), 'embeddings': bool(settings.openai_api_key), 'hh': bool(settings.hh_access_token),
        'hh_last': {'status': last.status, 'error': last.error, 'created_at': last.created_at.isoformat()} if last else None}


@app.post(PREFIX + '/documents')
def create_document(body: DocumentRequest, user=Depends(current_user), db=Depends(db_session), idempotency_key: str | None = Header(None)):
    owned(db, body.cv_id, user.id, 'cv')
    owned(db, body.vacancy_id, user.id, 'vacancy')
    return enqueue(db, user.id, 'document', body.model_dump(), request_key(idempotency_key))


@app.put(PREFIX + '/documents/{document_id}')
def edit_document(document_id: str, body: EditText, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, document_id, user.id, 'document')
    save_data(row, text=body.text, versions=row.data.get('versions', []) + [{'text': row.data['text'], 'saved_at': now().isoformat()}])
    row.status = 'saved'
    db.commit()
    return public_record(row)


@app.get(PREFIX + '/documents/{document_id}/export')
def export_document(document_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, document_id, user.id, 'document')
    document = Document()
    document.add_heading(row.data['title'], 0)
    for paragraph in row.data['text'].split('\n'):
        document.add_paragraph(paragraph)
    output = io.BytesIO()
    document.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                             headers={'Content-Disposition': 'attachment; filename="jobfinder-document.docx"'})


class PlanRequest(Strict):
    vacancy_id: str
    cv_id: str


@app.post(PREFIX + '/plans')
def plan_create(body: PlanRequest, user=Depends(current_user), db=Depends(db_session), idempotency_key: str | None = Header(None)):
    owned(db, body.cv_id, user.id, 'cv')
    owned(db, body.vacancy_id, user.id, 'vacancy')
    return enqueue(db, user.id, 'plan', {**body.model_dump(), 'profile': user.profile}, request_key(idempotency_key))


@app.post(PREFIX + '/plans/{plan_id}/days/{day}')
def plan_done(plan_id: str, day: int, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, plan_id, user.id, 'plan', lock=True)
    if not 1 <= day <= 7:
        raise HTTPException(422, 'День должен быть от 1 до 7')
    days = list(row.data['days'])
    days[day - 1] = {**days[day - 1], 'done': not days[day - 1]['done']}
    save_data(row, days=days)
    db.commit()
    return public_record(row)


@app.post(PREFIX + '/interviews')
def interview_create(body: InterviewInput, user=Depends(current_user), db=Depends(db_session)):
    vacancy = owned(db, body.vacancy_id, user.id, 'vacancy')
    rows = retrieve(db, vacancy.data['description'], body.direction, body.level, body.language, 100)
    questions = [r for r in rows if r.kind == 'question'][:5]
    if len(questions) < 5:
        raise HTTPException(409, 'Для интервью нужны 5 опубликованных вопросов выбранного направления, уровня и языка')
    turns = [{'question_id': q.id, 'question': q.data, 'rubric_version': q.data['version'],
              'materials': evidence(db, q.data), 'answer': '', 'evaluation': None} for q in questions]
    row = Record(kind='interview', owner_id=user.id, status='active', data={**body.model_dump(), 'turns': turns})
    db.add(row)
    db.commit()
    return public_record(row)


@app.post(PREFIX + '/interviews/{interview_id}/answers/{index}')
def answer(interview_id: str, index: int, body: AnswerInput, user=Depends(current_user), db=Depends(db_session)):
    row = owned(db, interview_id, user.id, 'interview', lock=True)
    if not 0 <= index < 5:
        raise HTTPException(422)
    if any(not turn.get('evaluation') for turn in row.data['turns'][:index]):
        raise HTTPException(409, 'Завершите предыдущий вопрос')
    if row.data['turns'][index].get('evaluation'):
        raise HTTPException(409, 'Ответ уже оценён')
    turns = list(row.data['turns'])
    turns[index] = {**turns[index], 'answer': body.text, 'submitted': True}
    save_data(row, turns=turns)
    payload = {'interview_id': row.id, 'index': index, 'text': body.text}
    return enqueue(db, user.id, 'evaluate', payload, f'answer:{row.id}:{index}')


@app.post(PREFIX + '/interviews/{interview_id}/audio')
async def answer_audio(interview_id: str, file: UploadFile = File(...), user=Depends(current_user), db=Depends(db_session)):
    owned(db, interview_id, user.id, 'interview')
    if not settings.openai_api_key:
        raise HTTPException(409, 'Распознавание голоса не подключено. Введите ответ текстом.')
    path = await upload(file, user.id, {'.webm', '.mp3', '.mp4', '.m4a', '.wav', '.ogg'}, 15_000_000)
    return enqueue(db, user.id, 'audio_answer', {'interview_id': interview_id, 'path': str(path)}, uid())


@app.get(PREFIX + '/stats')
def statistics(direction: Direction = 'frontend', level: Level = 'junior', user=Depends(current_user), db=Depends(db_session)):
    rows = db.scalars(select(Record).where(Record.owner_id == user.id, Record.kind == 'interview').order_by(Record.created_at)).all()
    timeline, topics, errors = [], defaultdict(list), defaultdict(int)
    for row in rows:
        if row.data['direction'] != direction or row.data['level'] != level:
            continue
        scores = []
        for turn in row.data['turns']:
            evaluation = turn.get('evaluation')
            if not evaluation or not evaluation['reliable']:
                continue
            score = sum(evaluation[k] for k in ('correctness', 'completeness', 'reasoning')) / 3
            scores.append(score)
            topics[turn['question']['topic']].append(score)
            for error in evaluation['errors']:
                errors[error] += 1
        if scores:
            timeline.append({'id': row.id, 'date': row.created_at.isoformat(), 'score': round(sum(scores) / len(scores), 2), 'answers': len(scores)})
    return {'timeline': timeline, 'topics': [{'topic': k, 'score': round(sum(v) / len(v), 2), 'answers': len(v)} for k, v in topics.items()],
            'errors': sorted([{'error': k, 'count': v} for k, v in errors.items()], key=lambda e: -e['count'])}


def job_dict(job):
    return {'id': job.id, 'kind': job.kind, 'status': job.status, 'result': job.result, 'error': job.error,
            'created_at': job.created_at.isoformat(), 'attempts': job.attempts, 'completed_steps': list(job.checkpoints)}


@app.get(PREFIX + '/jobs')
def job_list(user=Depends(current_user), db=Depends(db_session)):
    return [job_dict(j) for j in db.scalars(select(Job).where(Job.owner_id == user.id).order_by(Job.created_at.desc()).limit(100))]


@app.get(PREFIX + '/jobs/{job_id}')
def get_job(job_id: str, user=Depends(current_user), db=Depends(db_session)):
    job = db.get(Job, job_id)
    if not job or job.owner_id != user.id:
        raise HTTPException(404)
    return job_dict(job)


@app.post(PREFIX + '/jobs/{job_id}/resume')
def resume_job(job_id: str, user=Depends(current_user), db=Depends(db_session)):
    job = db.get(Job, job_id)
    if not job or job.owner_id != user.id:
        raise HTTPException(404)
    if job.status not in ('paused', 'failed') or job.attempts >= 3:
        raise HTTPException(409, 'Нельзя повторить: требуется проверка, достигнут лимит попыток или задание уже выполняется')
    job.status, job.error = 'queued', None
    db.commit()
    return job_dict(job)


@app.get(PREFIX + '/admin/{kind}')
def admin_list(kind: str, user=Depends(admin), db=Depends(db_session)):
    if kind not in ('source', 'question', 'material'):
        raise HTTPException(404)
    return [record_dict(r) for r in db.scalars(select(Record).where(Record.owner_id.is_(None), Record.kind == kind).order_by(Record.created_at.desc()))]


@app.post(PREFIX + '/admin/sources')
def source_create(body: SourceInput, user=Depends(admin), db=Depends(db_session)):
    video = youtube_id(body.url)
    row = db.scalar(select(Record).where(Record.kind == 'source', Record.owner_id.is_(None), Record.dedup_key == video))
    if row:
        if body.duration_seconds and not row.data.get('duration_seconds'):
            save_data(row, duration_seconds=body.duration_seconds)
            db.commit()
        return record_dict(row)
    row = Record(kind='source', dedup_key=video, data={**body.model_dump(), 'url': f'https://www.youtube.com/watch?v={video}', 'video_id': video})
    db.add(row)
    db.commit()
    return record_dict(row)


@app.post(PREFIX + '/admin/sources/{source_id}/import')
def source_import(source_id: str, force_audio: bool = False, user=Depends(admin), db=Depends(db_session)):
    source = shared(db, source_id, 'source')
    if source.data.get('transcript') and not force_audio:
        return {'record_id': source.id, 'message': 'Расшифровка уже сохранена'}
    return enqueue(db, user.id, 'import_source', {'source_id': source_id, 'force_audio': force_audio}, f'import:{source_id}:{force_audio}')


@app.post(PREFIX + '/admin/sources/{source_id}/upload')
async def source_upload(source_id: str, file: UploadFile = File(...), user=Depends(admin), db=Depends(db_session)):
    shared(db, source_id, 'source')
    path = await upload(file, user.id, {'.txt', '.srt', '.vtt', '.webm', '.mp3', '.mp4', '.m4a', '.wav', '.ogg'}, 150_000_000)
    return enqueue(db, user.id, 'import_source', {'source_id': source_id, 'path': str(path)}, uid())


@app.post(PREFIX + '/admin/sources/{source_id}/extract')
def source_extract(source_id: str, user=Depends(admin), db=Depends(db_session)):
    row = shared(db, source_id, 'source')
    version = hashlib.sha256(str(row.data.get('transcript')).encode()).hexdigest()[:16]
    return enqueue(db, user.id, 'extract_questions', {'source_id': source_id}, f'extract:{source_id}:{version}')


@app.post(PREFIX + '/admin/materials')
def material_create(body: MaterialInput, user=Depends(admin), db=Depends(db_session), idempotency_key: str | None = Header(None)):
    return enqueue(db, user.id, 'import_material', body.model_dump(), request_key(idempotency_key))


@app.post(PREFIX + '/admin/questions')
def question_create(body: QuestionInput, user=Depends(admin), db=Depends(db_session)):
    row = Record(kind='question', data={**body.model_dump(), 'sources': [], 'version': 1})
    db.add(row)
    db.commit()
    return record_dict(row)


@app.put(PREFIX + '/admin/questions/{question_id}')
def question_edit(question_id: str, body: QuestionInput, user=Depends(admin), db=Depends(db_session)):
    row = owned(db, question_id, None, 'question', lock=True)
    before = {k: v for k, v in row.data.items() if k != 'history'}
    row.data = {**row.data, **body.model_dump(), 'version': row.data.get('version', 1) + 1,
                'history': row.data.get('history', []) + [before]}
    if len(row.data.get('sources', [])) == 1:
        row.data = {**row.data, 'sources': [{**row.data['sources'][0], 'start': body.start, 'end': body.end}]}
    row.status = 'draft'
    db.execute(delete(Knowledge).where(Knowledge.record_id == row.id))
    db.commit()
    return record_dict(row)


@app.put(PREFIX + '/admin/materials/{material_id}')
def material_edit(material_id: str, body: EditText, user=Depends(admin), db=Depends(db_session)):
    row = owned(db, material_id, None, 'material', lock=True)
    save_data(row, text=body.text, version=row.data.get('version', 1) + 1)
    row.status = 'draft'
    db.execute(delete(Knowledge).where(Knowledge.record_id == row.id))
    db.commit()
    return record_dict(row)


@app.post(PREFIX + '/admin/questions/{question_id}/merge/{other_id}')
def merge_questions(question_id: str, other_id: str, user=Depends(admin), db=Depends(db_session)):
    if question_id == other_id:
        raise HTTPException(422)
    target, other = shared(db, question_id, 'question'), shared(db, other_id, 'question')
    sources = target.data.get('sources', []) + other.data.get('sources', [])
    unique = list({json_key(s): s for s in sources}.values())
    save_data(target, sources=unique, version=target.data.get('version', 1) + 1,
              merged_discussions=target.data.get('merged_discussions', []) + [other.data])
    target.status, other.status = 'draft', 'merged'
    db.execute(delete(Knowledge).where(Knowledge.record_id.in_([target.id, other.id])))
    db.commit()
    return record_dict(target)


def json_key(value):
    import json
    return json.dumps(value, sort_keys=True)


@app.post(PREFIX + '/admin/publish/{record_id}')
def publish(record_id: str, user=Depends(admin), db=Depends(db_session)):
    row = owned(db, record_id, None, lock=True)
    if row.kind not in ('question', 'material'):
        raise HTTPException(422)
    data = row.data
    if row.kind == 'question':
        q = QuestionInput.model_validate({k: v for k, v in data.items() if k in QuestionInput.model_fields})
        if q.needs_context or len(q.reference_answer) < 30 or len(q.rubric) < 3 or not q.material_ids:
            raise HTTPException(422, 'Для публикации дополните контекст, эталон, минимум 3 критерия и проверенные материалы')
        for material_id in q.material_ids:
            if shared(db, material_id, 'material').status != 'published':
                raise HTTPException(422, 'Сначала проверьте и опубликуйте материалы')
        if any(s.get('end', 0) <= s.get('start', 0) for s in data.get('sources', [])):
            raise HTTPException(422, 'Укажите точные таймкоды исходного обсуждения')
        content = q.question + '\n' + q.reference_answer + '\n' + ' '.join(q.rubric)
    else:
        content = data['text']
    row.status = 'published'
    save_data(row, reviewed_at=now().isoformat(), version=data.get('version', 1))
    index = db.get(Knowledge, row.id)
    if index and index.text != content:
        db.delete(index)
        db.flush()
        index = None
    if index is None:
        db.add(Knowledge(record_id=row.id, text=content, direction=data['direction'], level=data['level'], language=data['language']))
    db.commit()
    indexing = enqueue(db, user.id, 'index_knowledge', {'record_id': row.id, 'version': row.data['version']}, f'index:{row.id}:{row.data["version"]}:v2') if settings.openai_api_key else None
    return {'record': record_dict(row), 'indexing': indexing}


@app.get(PREFIX + '/admin/usage/summary')
def usage_summary(user=Depends(admin), db=Depends(db_session)):
    month = now().strftime('%Y-%m')
    budget = db.get(Budget, month)
    usage = db.scalars(select(Usage).where(Usage.month == month).order_by(Usage.created_at.desc())).all()
    return {'month': month, 'limit': settings.monthly_budget_usd, 'charged_and_reserved': float(budget.charged) if budget else 0,
        'video_hours': budget.video_seconds / 3600 if budget else 0,
        'operations': [{'key': u.key, 'model': u.model, 'state': u.state, 'reserved': float(u.reserved),
                        'actual': float(u.actual) if u.actual is not None else None} for u in usage]}
