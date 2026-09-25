import html
import http.client
import ipaddress
import json
import re
import socket
import ssl
import subprocess
import zipfile
from pathlib import Path
from urllib.parse import urlsplit, parse_qs, urljoin
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader
from youtube_transcript_api import YouTubeTranscriptApi
from .config import settings
from . import ai


def youtube_id(url):
    parts = urlsplit(url)
    if parts.scheme != 'https' or parts.username or parts.password or parts.port not in (None, 443):
        raise ValueError('Нужна HTTPS-ссылка YouTube')
    host = (parts.hostname or '').lower()
    if host == 'youtu.be':
        video = parts.path.strip('/')
    elif host in ('youtube.com', 'www.youtube.com', 'm.youtube.com'):
        if parts.path == '/watch':
            video = parse_qs(parts.query).get('v', [''])[0]
        elif parts.path.startswith(('/shorts/', '/embed/', '/live/')):
            video = parts.path.split('/')[2]
        else:
            video = ''
    else:
        video = ''
    if not re.fullmatch(r'[A-Za-z0-9_-]{11}', video):
        raise ValueError('Не удалось определить YouTube ID отдельного видео')
    return video


def run(args, timeout=180):
    result = subprocess.run(args, capture_output=True, text=True, timeout=timeout)
    if result.returncode:
        # Command stderr may contain signed media URLs. Do not expose them.
        raise ValueError(f'{Path(args[0]).name}: операция недоступна. Попробуйте ручную загрузку файла.')
    return result.stdout


def duration(path):
    result = json.loads(run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'json', str(path)]))
    value = result.get('format', {}).get('duration')
    if value not in (None, 'N/A'):
        seconds = float(value)
    else:
        # MediaRecorder WebM streams commonly omit the container duration.
        packets = json.loads(run(['ffprobe', '-v', 'error', '-select_streams', 'a:0', '-read_intervals', '%+10801',
            '-show_entries', 'packet=pts_time,duration_time', '-of', 'json', str(path)]))
        seconds = max((float(p.get('pts_time', 0)) + float(p.get('duration_time', 0)) for p in packets.get('packets', [])), default=0)
    if not 0 < seconds <= 10800:
        raise ValueError('Допустимая длительность: от 1 секунды до 3 часов')
    return seconds


def extract_cv(path):
    if path.suffix == '.pdf':
        reader = PdfReader(path)
        if len(reader.pages) > 40:
            raise ValueError('Не более 40 страниц в резюме')
        value = '\n'.join(page.extract_text() or '' for page in reader.pages)
    elif path.suffix == '.docx':
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30_000_000:
                raise ValueError('Слишком большой распакованный документ')
        doc = Document(path)
        value = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    else:
        raise ValueError('Поддерживаются PDF и DOCX')
    if len(value.strip()) < 40:
        raise ValueError('Текст не извлечён. Для скана вставьте текст резюме вручную.')
    if len(value) > 60000:
        raise ValueError('Не более 60 000 символов')
    return value


def timestamp(value):
    bits = value.replace(',', '.').split(':')
    return sum(float(v) * 60 ** i for i, v in enumerate(reversed(bits)))


def subtitles(raw):
    pattern = r'(?m)^(\d{1,2}:\d{2}(?::\d{2})?[.,]\d{3})\s+-->\s+(\d{1,2}:\d{2}(?::\d{2})?[.,]\d{3})[^\n]*\n(.*?)(?=\n\s*\n|\Z)'
    segments = []
    for match in re.finditer(pattern, raw.replace('\r', ''), re.S):
        text = html.unescape(re.sub(r'<[^>]*>', '', match[3])).strip()
        if text:
            segments.append({'start': timestamp(match[1]), 'end': timestamp(match[2]), 'text': text, 'speaker': None})
    return clean_segments(segments)


def clean_segments(segments):
    cleaned = []
    for item in sorted(segments, key=lambda s: s['start']):
        text = re.sub(r'\s+', ' ', item['text']).strip()
        if not text:
            continue
        # Remove rolling-caption overlap only for overlapping/adjacent cues.
        if cleaned and item['start'] <= cleaned[-1]['end'] + .2:
            previous = cleaned[-1]['text'].split()
            current = text.split()
            for size in range(min(len(previous), len(current)), 0, -1):
                if previous[-size:] == current[:size]:
                    text = ' '.join(current[size:])
                    break
        if text:
            cleaned.append({**item, 'text': text})
    return cleaned


def fetch_youtube(video, language, directory, force_audio=False):
    url = f'https://www.youtube.com/watch?v={video}'
    failures = []
    if not force_audio:
        try:
            tracks = list(YouTubeTranscriptApi().list(video))
            # Prefer authored tracks in requested/original language, then auto.
            tracks.sort(key=lambda t: (t.language_code.split('-')[0] != language, t.is_generated))
            if tracks:
                track = tracks[0]
                raw = track.fetch().to_raw_data()
                segments = [{'start': x['start'], 'end': x['start'] + x['duration'], 'text': x['text'], 'speaker': None} for x in raw]
                return {'raw': raw, 'segments': clean_segments(segments), 'method': 'youtube-auto' if track.is_generated else 'youtube-authored', 'language': track.language_code}
        except Exception as exc:
            failures.append(type(exc).__name__)
        try:
            run(['yt-dlp', '--no-playlist', '--skip-download', '--write-subs', '--write-auto-subs',
                 '--sub-langs', f'{language}.*,en.*,ru.*', '--sub-format', 'vtt', '--socket-timeout', '20',
                 '--retries', '1', '-o', str(directory / 'captions.%(ext)s'), url])
            files = sorted(directory.glob('*.vtt'), key=lambda p: language not in p.name)
            for file in files:
                raw = file.read_text(encoding='utf-8')
                segments = subtitles(raw)
                if segments:
                    return {'raw': raw, 'segments': segments, 'method': 'yt-dlp-subtitles', 'language': language}
        except Exception as exc:
            failures.append(type(exc).__name__)
    # Do not download hours of audio when transcription cannot run.
    if not settings.openai_api_key:
        raise ai.Paused('Субтитры недоступны. Загрузите TXT/SRT/VTT вручную или настройте OpenAI для аудио. ' + ', '.join(failures))
    run(['yt-dlp', '--no-playlist', '--match-filter', 'duration <= 10800 & !is_live', '-f', 'bestaudio',
         '--max-filesize', '300M', '--socket-timeout', '20', '--retries', '1',
         '-o', str(directory / 'audio.%(ext)s'), url], timeout=600)
    files = [p for p in directory.glob('audio.*') if p.suffix not in ('.part', '.ytdl')]
    if not files:
        raise ValueError('Аудио недоступно. Загрузите файл вручную.')
    return {'audio_path': str(files[0]), 'method': 'youtube-audio', 'language': language}


def audio_transcript(job_id, path, directory, diarize=True):
    seconds = duration(path)
    if not diarize and seconds > 181:
        raise ValueError('Ответ должен быть не длиннее 3 минут')
    if diarize:
        ai.reserve_video(job_id, seconds)
    segments, raw = [], []
    # 10-minute pieces at 48 kbps remain well under the 25 MB provider limit.
    # Two seconds of overlap avoid cutting words. clean_segments removes repeated
    # suffix/prefix words only where timestamps overlap. Speakers stay chunk-scoped.
    for index, start in enumerate(range(0, int(seconds) + 1, 598)):
        length = min(600, seconds - start)
        if length <= 0:
            continue
        chunk = directory / f'chunk-{index}.mp3'
        run(['ffmpeg', '-nostdin', '-v', 'error', '-y', '-ss', str(start), '-i', str(path), '-t', str(length),
             '-vn', '-ac', '1', '-ar', '16000', '-b:a', '48k', str(chunk)])
        result = ai.transcribe(job_id, f'audio-{index}', chunk, length, diarize)
        raw.append(result)
        parts = result.get('segments') or [{'start': 0, 'end': length, 'text': result['text']}]
        for part in parts:
            segments.append({'start': start + part['start'], 'end': start + part['end'], 'text': part['text'],
                             'speaker': f"chunk{index}:{part['speaker']}" if part.get('speaker') else None})
        chunk.unlink(missing_ok=True)
    return {'raw': raw, 'segments': clean_segments(segments), 'method': 'diarized-audio' if diarize else 'answer-audio', 'duration': seconds}


def public_page(url):
    """Allowlisted HTTPS, IP validation and pinned socket; validate each redirect."""
    for _ in range(4):
        parts = urlsplit(url)
        host = parts.hostname
        if parts.scheme != 'https' or parts.port not in (None, 443) or parts.username or parts.password or host not in settings.allowed_material_hosts.split(','):
            raise ValueError('URL должен быть HTTPS на домене из ALLOWED_MATERIAL_HOSTS')
        addresses = socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
            raise ValueError('Непубличный адрес запрещён')
        address = addresses[0][4][0]
        connection = http.client.HTTPSConnection(host, timeout=20)
        connection.sock = ssl.create_default_context().wrap_socket(socket.create_connection((address, 443), timeout=20), server_hostname=host)
        try:
            connection.request('GET', parts.path + ('?' + parts.query if parts.query else ''), headers={'User-Agent': 'JobFinderKZ/0.1'})
            response = connection.getresponse()
            if response.status in (301, 302, 303, 307, 308):
                url = urljoin(url, response.getheader('Location', ''))
                continue
            if response.status != 200 or 'text/html' not in response.getheader('Content-Type', ''):
                raise ValueError('Страница недоступна или не является HTML')
            content = response.read(2_000_001)
            if len(content) > 2_000_000:
                raise ValueError('Страница превышает 2 МБ')
            soup = BeautifulSoup(content, 'html.parser')
            for node in soup(['script', 'style', 'nav', 'footer', 'header']):
                node.decompose()
            main = soup.find('main') or soup.find('article') or soup
            return {'url': url, 'title': soup.title.get_text() if soup.title else host, 'text': main.get_text(' ', strip=True)[:60000]}
        finally:
            connection.close()
    raise ValueError('Слишком много перенаправлений')
