import io
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import pytest
from docx import Document
from pypdf import PdfWriter
from sqlalchemy import select
from app import ai, ingest, tasks
from app.config import settings
from app.db import Session, Job, Record, Budget, now, uid
from app.worker import run_once
from tests.conftest import register

P = '/api/v1'


def test_docx_upload_tables_original_and_delete(client, user):
    doc = Document()
    doc.add_paragraph('Разработчик Python. Учебные проекты и работа с PostgreSQL.')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Навыки: Python, SQL, pytest'
    output = io.BytesIO()
    doc.save(output)
    body = output.getvalue()
    response = client.post(P + '/cv/upload', files={'file': ('cv.docx', body)})
    assert response.status_code == 200, response.text
    row = response.json()
    assert 'pytest' in row['data']['text']
    assert 'file_path' not in row['data']
    assert client.get(P + f'/cv/{row["id"]}/original').content == body
    folder = settings.storage_path / 'users' / user['id']
    assert any(folder.iterdir())
    assert client.delete(P + '/account').status_code == 200
    assert not folder.exists()


def test_scan_and_corrupt_pdf_offer_manual_text(client, user):
    pdf = PdfWriter()
    pdf.add_blank_page(width=400, height=600)
    output = io.BytesIO()
    pdf.write(output)
    response = client.post(P + '/cv/upload', files={'file': ('scan.pdf', output.getvalue())})
    assert response.status_code == 422 and 'вручную' in response.json()['detail']
    response = client.post(P + '/cv/upload', files={'file': ('bad.pdf', b'invalid')})
    assert response.status_code == 422
    assert client.post(P + '/cv/upload', files={'file': ('bad.exe', b'invalid')}).status_code == 422
    assert not list((settings.storage_path / 'users' / user['id']).iterdir())


def test_two_workers_cannot_execute_same_job(client, user, monkeypatch):
    started, release = threading.Event(), threading.Event()
    calls = []
    def handler(job):
        calls.append(job.id)
        started.set()
        assert release.wait(5)
        return {'ok': True}
    monkeypatch.setitem(tasks.HANDLERS, 'lock-test', handler)
    with Session.begin() as db:
        job = Job(owner_id=user['id'], kind='lock-test', request_key=uid(), payload={})
        db.add(job)
    with ThreadPoolExecutor(max_workers=2) as pool:
        one = pool.submit(run_once)
        assert started.wait(3)
        two = pool.submit(run_once)
        try:
            assert two.result(timeout=3) is False
        finally:
            release.set()
        assert one.result(timeout=3) is True
    assert len(calls) == 1


def test_manual_txt_requires_duration_and_is_counted(client):
    register(client, 'owner@example.com')
    source = client.post(P + '/admin/sources', json={'url': 'https://youtu.be/GlK6nGzAK8E'}).json()
    response = client.post(P + f'/admin/sources/{source["id"]}/upload', files={'file': ('transcript.txt', 'Что такое транзакция? Изменения фиксируются вместе.'.encode())})
    assert run_once()
    job_id = response.json()['job_id']
    assert client.get(P + '/jobs/' + job_id).json()['status'] == 'failed'
    client.post(P + '/admin/sources', json={'url': 'https://youtu.be/GlK6nGzAK8E', 'duration_seconds': 120})
    assert client.post(P + f'/jobs/{job_id}/resume').status_code == 200
    assert run_once()
    assert client.get(P + '/jobs/' + job_id).json()['status'] == 'completed'
    with Session() as db:
        assert db.get(Budget, now().strftime('%Y-%m')).video_seconds == 120


def test_chunk_offsets_scoped_speakers_and_overlap(client, user, monkeypatch, tmp_path):
    monkeypatch.setattr(ingest, 'duration', lambda p: 700)
    monkeypatch.setattr(ingest, 'run', lambda *a, **k: '')
    def transcript(job_id, step, path, length, diarize):
        if step == 'audio-0':
            return {'segments': [{'start': 598, 'end': 600, 'text': 'hello world', 'speaker': 'A'}]}
        return {'segments': [{'start': 0, 'end': 3, 'text': 'hello world again', 'speaker': 'A'}]}
    monkeypatch.setattr(ai, 'transcribe', transcript)
    with Session.begin() as db:
        job = Job(owner_id=user['id'], kind='import_source', request_key=uid(), payload={})
        db.add(job)
        db.flush()
        job_id = job.id
    result = ingest.audio_transcript(job_id, tmp_path / 'input.mp3', tmp_path)
    assert [s['text'] for s in result['segments']] == ['hello world', 'again']
    assert [s['speaker'] for s in result['segments']] == ['chunk0:A', 'chunk1:A']
    assert result['segments'][1]['start'] == 598
    assert result['segments'][1]['end'] == 601


def test_invalid_ai_fact_id_does_not_create_document(client, user, monkeypatch):
    from app.schemas import DocumentResult
    from tests.test_system import cv, vacancy
    resume, role = cv(client), vacancy(client)
    monkeypatch.setattr(ai, 'structured', lambda *a, **k: DocumentResult(title='Bad', introduction='',
        selected_fact_ids=['invented:ten-years'], closing='', changes=[]))
    response = client.post(P + '/documents', json={'vacancy_id': role['id'], 'cv_id': resume['id'], 'kind': 'adapted_cv'})
    assert run_once()
    assert client.get(P + '/jobs/' + response.json()['job_id']).json()['status'] == 'failed'
    assert client.get(P + '/records/document').json() == []
